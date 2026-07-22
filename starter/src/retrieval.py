import re
import uuid
from dataclasses import dataclass
from typing import List, Dict, Any, Optional

import chromadb

from schemas import DocumentChunk

# Module-level cache: embedding_model id → {doc_id: (chunks, embeddings, metadatas)}
# Avoids re-calling the Gemini embedding API for sample docs on every new session.
_SAMPLE_EMBEDDING_CACHE: Dict[int, Dict[str, tuple]] = {}


@dataclass
class Document:
    """Represents a document in our system"""
    doc_id: str
    title: str
    content: str
    doc_type: str  # 'invoice', 'contract', 'claim'
    metadata: Dict[str, Any]


def _chunk_text(text: str, chunk_size: int = 500, overlap: int = 50) -> List[str]:
    """Split text into overlapping chunks."""
    chunks = []
    start = 0
    text = text.strip()
    while start < len(text):
        end = start + chunk_size
        chunks.append(text[start:end])
        start += chunk_size - overlap
    return chunks if chunks else [text]


def _extract_text(filename: str, file_bytes: bytes) -> str:
    """Extract plain text from PDF, DOCX, TXT, or CSV bytes."""
    ext = filename.rsplit(".", 1)[-1].lower()

    if ext == "pdf":
        import io
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(file_bytes))
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    if ext == "docx":
        import io
        from docx import Document as DocxDocument
        doc = DocxDocument(io.BytesIO(file_bytes))
        return "\n".join(p.text for p in doc.paragraphs)

    # txt / csv / anything else — treat as utf-8 text
    return file_bytes.decode("utf-8", errors="replace")


class ChromaRetriever:
    """
    Per-session ChromaDB in-memory retriever backed by Gemini (or OpenAI) embeddings.
    Each instance owns one named collection so users are fully isolated.
    """

    def __init__(self, embedding_model, collection_name: str):
        self.embedding_model = embedding_model
        self.client = chromadb.EphemeralClient()
        self.collection = self.client.get_or_create_collection(
            name=collection_name,
            metadata={"hnsw:space": "cosine"},
        )
        # Track uploaded filenames for the UI
        self.ingested_files: List[str] = []
        # Keep a lightweight doc registry for statistics / amount queries
        self._doc_registry: Dict[str, Document] = {}
        self._load_sample_documents()

    # ------------------------------------------------------------------
    # Ingestion
    # ------------------------------------------------------------------

    def ingest_file(self, filename: str, file_bytes: bytes) -> int:
        """
        Parse, chunk, embed, and store a file in the collection.
        Returns number of chunks stored.
        """
        text = _extract_text(filename, file_bytes)
        chunks = _chunk_text(text)
        if not chunks:
            return 0

        doc_id = filename.rsplit(".", 1)[0]
        ext = filename.rsplit(".", 1)[-1].lower()
        doc_type = "pdf" if ext == "pdf" else ("docx" if ext == "docx" else "text")

        embeddings = self.embedding_model.embed_documents(chunks)

        ids = [f"{doc_id}_chunk_{i}_{uuid.uuid4().hex[:6]}" for i in range(len(chunks))]
        metadatas = [
            {"doc_id": doc_id, "filename": filename, "chunk_index": i, "doc_type": doc_type}
            for i in range(len(chunks))
        ]

        self.collection.add(
            ids=ids,
            embeddings=embeddings,
            documents=chunks,
            metadatas=metadatas,
        )

        if filename not in self.ingested_files:
            self.ingested_files.append(filename)

        # Register in doc_registry so statistics / id lookup works
        self._doc_registry[doc_id] = Document(
            doc_id=doc_id,
            title=filename,
            content=text,
            doc_type=doc_type,
            metadata={"filename": filename},
        )

        return len(chunks)

    def _ingest_document(self, doc: Document) -> None:
        """Ingest a Document dataclass directly (used for sample docs)."""
        model_key = id(self.embedding_model)
        cache = _SAMPLE_EMBEDDING_CACHE.setdefault(model_key, {})

        if doc.doc_id not in cache:
            chunks = _chunk_text(doc.content)
            embeddings = self.embedding_model.embed_documents(chunks)
            metadatas = [
                {
                    "doc_id": doc.doc_id,
                    "filename": doc.title,
                    "chunk_index": i,
                    "doc_type": doc.doc_type,
                    **{k: str(v) for k, v in doc.metadata.items()},
                }
                for i in range(len(chunks))
            ]
            cache[doc.doc_id] = (chunks, embeddings, metadatas)

        chunks, embeddings, metadatas = cache[doc.doc_id]
        ids = [f"{doc.doc_id}_chunk_{i}" for i in range(len(chunks))]

        self.collection.add(
            ids=ids,
            embeddings=embeddings,
            documents=chunks,
            metadatas=metadatas,
        )
        self._doc_registry[doc.doc_id] = doc

    def _load_sample_documents(self) -> None:
        """Load the 5 built-in sample documents."""
        samples = [
            Document(
                doc_id="INV-001",
                title="Invoice #12345",
                content="""
                Invoice #12345
                Date: 2024-01-15
                Client: Acme Corporation

                Services Rendered:
                - Consulting Services: $5,000
                - Software Development: $12,500
                - Support & Maintenance: $2,500

                Subtotal: $20,000
                Tax (10%): $2,000
                Total Due: $22,000

                Payment Terms: Net 30 days
                """,
                doc_type="invoice",
                metadata={"client": "Acme Corporation", "date": "2024-01-15", "total": "22000"},
            ),
            Document(
                doc_id="CON-001",
                title="Service Agreement",
                content="""
                SERVICE AGREEMENT

                This Service Agreement is entered into on January 1, 2024, between:
                - Provider: DocDacity Solutions Inc.
                - Client: Healthcare Partners LLC

                Services:
                1. Document Processing Platform Access
                2. 24/7 Technical Support
                3. Monthly Data Analytics Reports
                4. Compliance Monitoring

                Duration: 12 months
                Monthly Fee: $15,000
                Total Contract Value: $180,000

                Termination: Either party may terminate with 60 days written notice.
                """,
                doc_type="contract",
                metadata={"value": "180000", "duration_months": "12", "client": "Healthcare Partners LLC"},
            ),
            Document(
                doc_id="CLM-001",
                title="Insurance Claim #78901",
                content="""
                INSURANCE CLAIM FORM
                Claim Number: 78901
                Date of Incident: 2024-02-10
                Policy Number: POL-456789

                Claimant: John Doe
                Type of Claim: Medical Expense Reimbursement

                Expenses:
                - Hospital Visit: $1,200
                - Diagnostic Tests: $800
                - Medication: $150
                - Follow-up Consultation: $300

                Total Claim Amount: $2,450

                Status: Under Review
                """,
                doc_type="claim",
                metadata={"amount": "2450", "status": "Under Review", "claimant": "John Doe"},
            ),
            Document(
                doc_id="INV-002",
                title="Invoice #12346",
                content="""
                Invoice #12346
                Date: 2024-02-20
                Client: TechStart Inc.

                Products:
                - Enterprise License (Annual): $50,000
                - Implementation Services: $15,000
                - Training Package: $5,000

                Subtotal: $70,000
                Discount (10%): -$7,000
                Tax (10%): $6,300
                Total Due: $69,300

                Payment Terms: Net 45 days
                """,
                doc_type="invoice",
                metadata={"total": "69300", "client": "TechStart Inc.", "date": "2024-02-20"},
            ),
            Document(
                doc_id="INV-003",
                title="Invoice #12347",
                content="""
                Invoice #12347
                Date: 2024-03-01
                Client: Global Corp

                Services:
                - Annual Subscription: $120,000
                - Premium Support: $30,000
                - Custom Development: $45,000

                Subtotal: $195,000
                Tax (10%): $19,500
                Total Due: $214,500

                Payment Terms: Net 60 days
                """,
                doc_type="invoice",
                metadata={"total": "214500", "client": "Global Corp", "date": "2024-03-01"},
            ),
        ]
        for doc in samples:
            self._ingest_document(doc)

    # ------------------------------------------------------------------
    # Retrieval helpers
    # ------------------------------------------------------------------

    def _query_to_chunks(self, results: dict) -> List[DocumentChunk]:
        """Convert a chromadb query result dict to DocumentChunk list."""
        chunks = []
        docs = results.get("documents", [[]])[0]
        metas = results.get("metadatas", [[]])[0]
        distances = results.get("distances", [[]])[0]
        for doc, meta, dist in zip(docs, metas, distances):
            score = max(0.0, 1.0 - dist)
            chunks.append(DocumentChunk(
                doc_id=meta.get("doc_id", "unknown"),
                content=doc,
                metadata=meta,
                relevance_score=score,
            ))
        return chunks

    def _get_to_chunks(self, results: dict) -> List[DocumentChunk]:
        """Convert a chromadb get() result dict to DocumentChunk list."""
        chunks = []
        docs = results.get("documents") or []
        metas = results.get("metadatas") or []
        for doc, meta in zip(docs, metas):
            chunks.append(DocumentChunk(
                doc_id=meta.get("doc_id", "unknown"),
                content=doc,
                metadata=meta,
                relevance_score=1.0,
            ))
        return chunks

    # ------------------------------------------------------------------
    # Public retrieval interface (matches old SimulatedRetriever API)
    # ------------------------------------------------------------------

    def retrieve_by_keyword(self, query: str, top_k: int = 5) -> List[DocumentChunk]:
        embedding = self.embedding_model.embed_query(query)
        n = min(top_k, self.collection.count())
        if n == 0:
            return []
        results = self.collection.query(
            query_embeddings=[embedding],
            n_results=n,
        )
        return self._query_to_chunks(results)

    def retrieve_all(self) -> List[DocumentChunk]:
        results = self.collection.get(include=["documents", "metadatas"])
        return self._get_to_chunks(results)

    def retrieve_by_type(self, doc_type: str) -> List[DocumentChunk]:
        results = self.collection.get(
            where={"doc_type": doc_type},
            include=["documents", "metadatas"],
        )
        return self._get_to_chunks(results)

    def get_document_by_id(self, doc_id: str) -> Optional[DocumentChunk]:
        results = self.collection.get(
            where={"doc_id": doc_id},
            include=["documents", "metadatas"],
        )
        chunks = self._get_to_chunks(results)
        if not chunks:
            return None
        # Merge all chunks for the doc into one DocumentChunk
        full_content = "\n".join(c.content for c in chunks)
        return DocumentChunk(
            doc_id=doc_id,
            content=full_content,
            metadata=chunks[0].metadata,
            relevance_score=1.0,
        )

    def _get_document_amount(self, doc_id: str) -> Optional[float]:
        """Read amount from the doc registry metadata."""
        doc = self._doc_registry.get(doc_id)
        if not doc:
            return None
        for field in ("total", "amount", "value", "total_amount", "total_value"):
            val = doc.metadata.get(field)
            if val is not None:
                try:
                    return float(val)
                except (ValueError, TypeError):
                    continue
        return None

    def _chunks_with_amounts(self) -> List[DocumentChunk]:
        """Return one DocumentChunk per doc that has an amount in metadata."""
        seen = set()
        result = []
        for doc_id, doc in self._doc_registry.items():
            if doc_id in seen:
                continue
            if self._get_document_amount(doc_id) is not None:
                seen.add(doc_id)
                result.append(DocumentChunk(
                    doc_id=doc.doc_id,
                    content=doc.content,
                    metadata={**doc.metadata, "doc_type": doc.doc_type, "title": doc.title},
                    relevance_score=1.0,
                ))
        return result

    def retrieve_by_amount_range(
        self,
        min_amount: Optional[float] = None,
        max_amount: Optional[float] = None,
    ) -> List[DocumentChunk]:
        candidates = self._chunks_with_amounts()
        result = []
        for chunk in candidates:
            amount = self._get_document_amount(chunk.doc_id)
            if amount is None:
                continue
            if min_amount is not None and amount < min_amount:
                continue
            if max_amount is not None and amount > max_amount:
                continue
            result.append(chunk)
        result.sort(key=lambda c: self._get_document_amount(c.doc_id) or 0, reverse=True)
        return result

    def retrieve_by_exact_amount(self, amount: float, tolerance: float = 0.01) -> List[DocumentChunk]:
        return [
            c for c in self._chunks_with_amounts()
            if self._get_document_amount(c.doc_id) is not None
            and abs((self._get_document_amount(c.doc_id) or 0) - amount) <= tolerance
        ]

    def retrieve_by_approximate_amount(self, amount: float, percentage: float = 10.0) -> List[DocumentChunk]:
        tol = amount * (percentage / 100)
        result = []
        for chunk in self._chunks_with_amounts():
            doc_amount = self._get_document_amount(chunk.doc_id)
            if doc_amount is not None and abs(doc_amount - amount) <= tol:
                distance = abs(doc_amount - amount)
                chunk.relevance_score = 1.0 - (distance / tol)
                result.append(chunk)
        result.sort(key=lambda c: c.relevance_score, reverse=True)
        return result

    def _parse_and_retrieve_by_amount(self, query: str) -> List[DocumentChunk]:
        query_lower = query.lower()
        amount_pattern = r'\$?(\d+(?:,\d{3})*(?:\.\d{2})?)'
        amounts = [
            float(m.replace(",", "").replace("$", ""))
            for m in re.findall(amount_pattern, query)
        ]

        if any(w in query_lower for w in ["over", "above", "more than", "greater than", ">"]):
            if amounts:
                return self.retrieve_by_amount_range(min_amount=amounts[0])
        elif any(w in query_lower for w in ["under", "below", "less than", "<"]):
            if amounts:
                return self.retrieve_by_amount_range(max_amount=amounts[0])
        elif any(w in query_lower for w in ["between", "range", "from"]):
            if len(amounts) >= 2:
                return self.retrieve_by_amount_range(min_amount=min(amounts[:2]), max_amount=max(amounts[:2]))
        elif any(w in query_lower for w in ["around", "about", "approximately", "roughly", "~"]):
            if amounts:
                return self.retrieve_by_approximate_amount(amounts[0])
        elif any(w in query_lower for w in ["exactly", "exact", "precisely", "="]):
            if amounts:
                return self.retrieve_by_exact_amount(amounts[0])

        if amounts:
            return self.retrieve_by_amount_range(
                min_amount=min(amounts) * 0.9,
                max_amount=max(amounts) * 1.1,
            )
        return self.retrieve_by_keyword(query)

    def get_statistics(self) -> Dict[str, Any]:
        total_docs = len(self._doc_registry)
        doc_types: Dict[str, int] = {}
        amounts = []

        for doc_id, doc in self._doc_registry.items():
            doc_types[doc.doc_type] = doc_types.get(doc.doc_type, 0) + 1
            amount = self._get_document_amount(doc_id)
            if amount is not None:
                amounts.append(amount)

        stats: Dict[str, Any] = {
            "total_documents": total_docs,
            "documents_with_amounts": len(amounts),
            "total_amount": sum(amounts),
            "average_amount": sum(amounts) / len(amounts) if amounts else 0,
            "document_types": doc_types,
        }
        if amounts:
            stats["min_amount"] = min(amounts)
            stats["max_amount"] = max(amounts)
        return stats

    # Keep old property so list_documents() in main.py still works
    @property
    def documents(self) -> Dict[str, Document]:
        return self._doc_registry
